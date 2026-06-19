import re
from typing import List, Tuple, Optional, Dict

from openpyxl.worksheet.worksheet import Worksheet

THICK_STYLES = {"medium", "thick", "double"}


def _build_merged_lookup(ws: Worksheet) -> dict:
    lookup = {}
    for mrange in ws.merged_cells.ranges:
        tl = mrange.min_row, mrange.min_col
        for r in range(mrange.min_row, mrange.max_row + 1):
            for c in range(mrange.min_col, mrange.max_col + 1):
                lookup[(r, c)] = tl
    return lookup


def _cell_value(ws: Worksheet, r: int, c: int, merged_lookup: dict):
    tl = merged_lookup.get((r, c))
    if tl is not None:
        r, c = tl
    val = ws.cell(r, c).value
    return _norm(val)


def _norm(s) -> Optional[str]:
    if s is None:
        return None
    if not isinstance(s, str):
        return str(s)
    s = s.replace("\u3000", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return s


def _row_has_any_value(ws: Worksheet, merged_lookup: dict, r: int, max_col: int) -> bool:
    for c in range(1, max_col + 1):
        if _cell_value(ws, r, c, merged_lookup) not in (None, ""):
            return True
    return False


def _is_top_left(merged_lookup: dict, r: int, c: int) -> bool:
    return merged_lookup.get((r, c), (r, c)) == (r, c)


def _first_right_value(ws: Worksheet, merged_lookup: dict, max_col: int, r: int, start_col: int = 1) -> Optional[
    Tuple[int, int, str]]:
    c = start_col + 1
    while c <= max_col:
        if _is_top_left(merged_lookup, r, c):
            v = _cell_value(ws, r, c, merged_lookup)
            if v not in (None, ""):
                return r, c, str(v)
        c += 1
    return None


def _first_down_value(ws: Worksheet, merged_lookup: dict, max_row: int, max_col: int, start_row: int = 1,
                      start_col: int = 1) -> Optional[Tuple[int, int, str]]:
    r = start_row + 1
    while r <= max_row:
        for c in range(start_col, max_col + 1):
            if not _is_top_left(merged_lookup, r, c):
                continue
            v = _cell_value(ws, r, c, merged_lookup)
            if v not in (None, ""):
                return r, c, str(v)
        r += 1
    return None


def _cell_border_style_name(side) -> str:
    return (getattr(side, "style", None) or "") or ""


def _row_has_thick_horizontal_line(ws, merged_lookup, max_col, rr: int) -> bool:
    # 上一行 bottom
    if rr > 1:
        for cc in range(1, max_col + 1):
            r0, c0 = merged_lookup.get((rr - 1, cc), (rr - 1, cc))
            b = ws.cell(r0, c0).border
            if _cell_border_style_name(b.bottom) in THICK_STYLES:
                return True
    # 本行 top
    for cc in range(1, max_col + 1):
        r0, c0 = merged_lookup.get((rr, cc), (rr, cc))
        b = ws.cell(r0, c0).border
        if _cell_border_style_name(b.top) in THICK_STYLES:
            return True
    return False


def extract_section_usecase(ws, sheet_name, target_headers, merged_lookup, max_row, max_col):
    if not target_headers:
        return []

    buckets: Dict[str, List[Tuple[str, str]]] = {h: [] for h in target_headers}

    for r in range(1, max_row + 1):
        current_header = _cell_value(ws, r, 1, merged_lookup)

        if current_header not in target_headers:
            continue

        key = val = None
        c = 2
        while c <= max_col:
            v = _cell_value(ws, r, c, merged_lookup)
            if (v if not isinstance(v, str) else v.strip()) not in (None, ""):
                key = str(v).strip()
                c += 1
                break
            c += 1
        if key is None:
            continue

        while c <= max_col:
            v = _cell_value(ws, r, c, merged_lookup)
            if (v if not isinstance(v, str) else v.strip()) not in (None, ""):
                val = str(v).strip()
                break
            c += 1
        if val is None:
            continue

        buckets[current_header].append((key, val))

    buckets = {k: v for k, v in buckets.items() if v}
    result = [f"ユースケース情報 ——{sheet_name}"]
    result.extend([
        ("; ".join([f"{h}.{k}：{v}" for k, v in buckets[h]]) if buckets[h] else "") + f" ——{sheet_name}"
        for h in buckets.keys()
    ])
    return result


def extract_section_table(ws, sheet_name, target_tables, merged_lookup, max_row, max_col):
    if isinstance(target_tables, str):
        target_tables = [target_tables]

    results_all = []

    for target_table in target_tables:
        start_row = None
        for r in range(1, max_row + 1):
            v = _cell_value(ws, r, 1, merged_lookup)
            if v == target_table:
                start_row = r
                break
        if start_row is None:
            continue

        title_parts = []
        for c in range(1, max_col + 1):
            if not _is_top_left(merged_lookup, start_row, c):
                continue
            v = _cell_value(ws, start_row, c, merged_lookup)
            if v not in (None, ""):
                title_parts.append(str(v))
        table_title = "; ".join(title_parts) + f" ——{sheet_name}"

        header_row = None
        r = start_row + 1
        while r <= max_row:
            if _row_has_any_value(ws, merged_lookup, r, max_col):
                header_row = r
                break
            r += 1
        if header_row is None:
            results_all.append(table_title)
            continue

        headers = []
        for c in range(1, max_col + 1):
            h = _cell_value(ws, header_row, c, merged_lookup)
            headers.append(h if h not in ("", None) else None)

        results = [table_title]
        r = header_row + 1
        while r <= max_row:
            if not _row_has_any_value(ws, merged_lookup, r, max_col):
                break  # 到达该段落末尾

            parts = []
            seen_pairs = set()
            for c in range(1, max_col + 1):
                if not _is_top_left(merged_lookup, r, c):
                    continue
                head = headers[c - 1]
                if head is None:
                    continue
                val = _cell_value(ws, r, c, merged_lookup)
                if val in (None, ""):
                    continue

                pair = (str(head), str(val))
                if pair in seen_pairs:
                    continue
                seen_pairs.add(pair)

                parts.append(f"{head}：{val}")

            if parts:
                results.append("; ".join(parts) + f" ——{sheet_name}")
            r += 1

        results_all.extend(results)
    return results_all


def extract_section_shorirojikku(ws, sheet_name, merged_lookup, max_row, max_col, empty_row_stop=3):
    # 1) 找到锚点“処理ロジック”（第1列）
    start_row = None
    for r in range(1, max_row + 1):
        v = _cell_value(ws, r, 1, merged_lookup)
        if v == "処理ロジック":
            start_row = r
            break
    if start_row is None:
        return []

    # 2) 从下一行开始读取，直到停止条件
    lines = [f"処理ロジック ——{sheet_name}"]
    empty_streak = 0
    r = start_row + 1
    while r <= max_row:
        # 停止条件：见到“返却値”
        first_col = _cell_value(ws, r, 1, merged_lookup)
        if first_col in ("処理ロジック詳細", "処理フロー"):
            lines.append(f"{first_col} ——{sheet_name}")
            r += 1
            continue

        if first_col == "返却値":
            break

        # 停止条件：较粗横线
        if _row_has_thick_horizontal_line(ws, merged_lookup, max_col, r):
            break

        # 连续空行计数
        if not _row_has_any_value(ws, merged_lookup, r, max_col):
            empty_streak += 1
            if empty_streak >= empty_row_stop:
                break
            r += 1
            continue
        else:
            empty_streak = 0

        # 3) 组装该行字符串
        parts = []
        for c in range(1, max_col + 1):
            v = _cell_value(ws, r, c, merged_lookup)
            if v not in (None, ""):
                parts.append(str(v))
        if parts:
            lines.append("; ".join(parts) + f" ——{sheet_name}")

        r += 1

    return lines


def extract_section_teigi(ws, sheet_name, merged_lookup, max_row, max_col):
    # 1) 定位唯一 “SQL定義” 起始行
    start_row = None
    for r in range(1, max_row + 1):
        if _cell_value(ws, r, 1, merged_lookup) == "SQL定義":
            start_row = r
            break
    if start_row is None:
        return []

    # 2) 组装分段标题（把起始行上所有合并块左上的非空文本直接顺序拼接，适配 “SQL定義(SQL-01)” 等）
    title_parts = []
    for c in range(1, max_col + 1):
        if not _is_top_left(merged_lookup, start_row, c):
            continue
        v = _cell_value(ws, start_row, c, merged_lookup)
        if v not in (None, ""):
            title_parts.append(str(v))

    results = []
    results.append(f"SQL定義 ——{sheet_name}")

    # 3) 收集四个键的值（从 start_row 往下扫描整表）
    keys = ["SQL-ID", "SQL論理名", "対象テーブル", "操作"]
    kv = {}

    for r in range(start_row + 1, max_row + 1):
        k = _cell_value(ws, r, 1, merged_lookup)
        if k in keys and k not in kv:
            val = _first_right_value(ws, merged_lookup, max_col, r, start_col=1)[2]
            if val not in (None, ""):
                kv[k] = val

    # 4) 输出四项拼接行（仅输出拿到的项，顺序固定为 keys）
    parts = []
    for k in keys:
        if k in kv:
            parts.append(f"{k}：{kv[k]}")
    if parts:
        results.append("; ".join(parts) + f"  ——{sheet_name}")

    # 5) “使用目的”：先找到该行，再向下取首个有效值
    purpose_row = None
    purpose_col = None
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if _cell_value(ws, r, c, merged_lookup) == "使用目的":
                purpose_row = r
                purpose_col = c
                break
    if purpose_row is not None:
        purpose_val = _first_down_value(ws, merged_lookup, max_row, max_col, purpose_row, purpose_col)[2]
        if purpose_val not in (None, ""):
            results.append(f"使用目的：{purpose_val}  ——{sheet_name}")

    return results


def extract_section_sql(ws, sheet_name, merged_lookup, max_row, max_col):
    # -------- ① 寻找表头行，并确定 “SQL” 与 “備考” 的宽度（同名连续列 = 宽度） --------
    header_row = None
    sql_range = None  # (c_start, c_end)
    biko_range = None  # (c_start, c_end)
    # 也收集所有表头 -> 列区间，便于 “其它内容”
    header_ranges = {}  # {header_text: (c_start, c_end)}

    def scan_header_ranges(r: int):
        ranges = {}
        c = 1
        while c <= max_col:
            hv = _cell_value(ws, r, c, merged_lookup)
            if hv in (None, ""):
                c += 1
                continue
            # 连续相同表头名的区间
            start = c
            while c + 1 <= max_col and _cell_value(ws, r, c + 1, merged_lookup) == hv:
                c += 1
            end = c
            ranges[str(hv)] = (start, end)
            c += 1
        return ranges

    for r in range(1, max_row + 1):
        # 找到同时包含 “SQL” 与 “備考” 的行作为表头
        row_vals = [(_cell_value(ws, r, c, merged_lookup) or "") for c in range(1, max_col + 1)]
        if ("SQL" in row_vals) and ("備考" in row_vals):
            header_ranges = scan_header_ranges(r)
            if "SQL" in header_ranges and "備考" in header_ranges:
                header_row = r
                sql_range = header_ranges["SQL"]
                biko_range = header_ranges["備考"]
                break

    if header_row is None:
        return []  # 未找到有效表头则无内容

    # -------- ② 定义“按与key同宽取值并连接”的工具 --------
    def join_range_values(row: int, c_start: int, c_end: int, remove_spaces: bool) -> str:
        parts = []
        for c in range(c_start, c_end + 1):
            if _is_top_left(merged_lookup, row, c):
                v = _cell_value(ws, row, c, merged_lookup)
                if v not in (None, ""):
                    s = str(v).strip()
                    if remove_spaces:
                        # 按需求：去除空格后连接
                        s = s.replace(" ", "")
                    parts.append(s)
        # 不额外添加分隔，直接拼接
        return "".join(parts)

    # -------- ③~⑤ 逐行抽取，直到停止条件，组装每行字符串并汇总 --------
    results = []
    empty_streak = 0
    r = header_row + 1
    if merged_lookup[(r, header_ranges["SQL"][0])] == (header_row, header_ranges["SQL"][0]):
        r += 1

    while r <= max_row:
        # 停止：遇到“6.特記事項”
        first_col_val = _cell_value(ws, r, 1, merged_lookup)
        if first_col_val == "6.特記事項":
            break

        # 停止：较粗横线
        if _row_has_thick_horizontal_line(ws, merged_lookup, max_col, r):
            break

        # 空行计数
        if not _row_has_any_value(ws, merged_lookup, r, max_col):
            empty_streak += 1
            if empty_streak >= 2:  # 连续空行阈值=2
                break
            r += 1
            continue
        else:
            empty_streak = 0

        parts = []

        # 先取 SQL / 備考
        if sql_range:
            sql_val = join_range_values(r, sql_range[0], sql_range[1], remove_spaces=True)
            if sql_val:
                parts.append(f"SQL：{sql_val}")

        if biko_range:
            biko_val = join_range_values(r, biko_range[0], biko_range[1], remove_spaces=True)
            if biko_val:
                parts.append(f"備考：{biko_val}")

        # 其它可能存在的列（排除 SQL/備考）
        for h, (c0, c1) in header_ranges.items():
            if h in ("SQL", "備考"):
                continue
            val = join_range_values(r, c0, c1, remove_spaces=False)
            if val not in ("", None):
                parts.append(f"{h}：{val}")

        if parts:
            results.append("; ".join(parts) + f" ——{sheet_name}")

        r += 1

    if results:
        return [f"5.論理SQL ——{sheet_name}"] + results
    else:
        return []


def extract_sheet_shori(ws: Worksheet, sheet_name: str) -> List[str]:
    lines = []
    merged_lookup = _build_merged_lookup(ws)
    max_row, max_col = ws.max_row, ws.max_column
    target_headers = ["業務", "ユースケース", "パッケージ名", "処理"]
    target_tables = ["グローバル変数"]
    lines.extend(extract_section_usecase(ws, sheet_name, target_headers, merged_lookup, max_row, max_col))
    lines.extend(extract_section_shorirojikku(ws, sheet_name, merged_lookup, max_row, max_col))
    lines.extend(extract_section_table(ws, sheet_name, target_tables, merged_lookup, max_row, max_col))
    return lines


def extract_sheet_jikkoseigyo(ws: Worksheet, sheet_name: str) -> List[str]:
    lines = []
    merged_lookup = _build_merged_lookup(ws)
    max_row, max_col = ws.max_row, ws.max_column
    target_headers = ["パッケージ名", "処理"]
    target_tables = ["実行ロジック"]
    lines.extend(extract_section_usecase(ws, sheet_name, target_headers, merged_lookup, max_row, max_col))
    lines.extend(extract_section_table(ws, sheet_name, target_tables, merged_lookup, max_row, max_col))
    return lines


def extract_sheet_rojikku(ws: Worksheet, sheet_name: str) -> List[str]:
    lines = []
    merged_lookup = _build_merged_lookup(ws)
    max_row, max_col = ws.max_row, ws.max_column
    target_headers = ["ロジック", "ファンクション"]
    target_tables = ["パラメータ", "返却値"]
    lines.extend(extract_section_usecase(ws, sheet_name, target_headers, merged_lookup, max_row, max_col))
    lines.extend(extract_section_table(ws, sheet_name, target_tables, merged_lookup, max_row, max_col))
    lines.extend(extract_section_shorirojikku(ws, sheet_name, merged_lookup, max_row, max_col))
    # lines.extend(extract_section_table(ws, sheet_name, "返却値", merged_lookup, max_row, max_col))
    return lines


def extract_sheet_sql_teigi(ws: Worksheet, sheet_name: str) -> List[str]:
    lines = []
    merged_lookup = _build_merged_lookup(ws)
    max_row, max_col = ws.max_row, ws.max_column
    target_tables = ["1.入力パラメータ", "2.設定項目", "3.取得項目", "4.別名定義"]
    lines.extend(extract_section_teigi(ws, sheet_name, merged_lookup, max_row, max_col))
    lines.extend(extract_section_table(ws, sheet_name, target_tables, merged_lookup, max_row, max_col))
    lines.extend(extract_section_sql(ws, sheet_name, merged_lookup, max_row, max_col))
    target_tables = ["6.特記事項"]
    lines.extend(extract_section_table(ws, sheet_name, target_tables, merged_lookup, max_row, max_col))
    return lines
