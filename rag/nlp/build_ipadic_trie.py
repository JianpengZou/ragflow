import argparse
import csv
import math
import re
import string
from collections.abc import Iterable
from pathlib import Path

try:
    import datrie
except ImportError:
    datrie = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

DENOMINATOR = 1_000_000


def _encode_key(text: str) -> str:
    return str(text.lower().encode("utf-8"))[2:-1]


def _encode_rkey(text: str) -> str:
    return str(("DD" + (text[::-1].lower())).encode("utf-8"))[2:-1]


def _derive_pos_tag(row: list[str]) -> str:
    parts = [p for p in row[4:8] if p and p != "*"]
    return "-".join(parts) if parts else "UNKNOWN"


def _compute_frequency(cost: int) -> int:
    return max(1, int(DENOMINATOR / (cost + 1)))


def load_entries(csv_files: Iterable[Path]) -> dict[str, tuple[int, str]]:
    entries: dict[str, tuple[int, str]] = {}
    for csv_path in csv_files:
        with csv_path.open("r", encoding="euc-jp", errors="ignore", newline="") as fd:
            reader = csv.reader(fd)
            for row in reader:
                if len(row) < 5:
                    continue
                surface = row[0].strip()
                if not surface:
                    continue
                try:
                    cost = int(row[3])
                except ValueError:
                    continue
                freq = _compute_frequency(cost)
                pos_tag = _derive_pos_tag(row)
                if surface not in entries or freq > entries[surface][0]:
                    entries[surface] = (freq, pos_tag)
    return entries


def _normalize_category(category: str | None) -> str:
    if not category:
        return "名詞-IT用語"
    slug = re.sub(r"\s+", "", str(category))
    return f"名詞-IT用語-{slug}"


def _extract_terms(raw: str | None) -> Iterable[str]:
    if not raw:
        return []
    text = str(raw).replace("／", "|")
    candidates: list[str] = []
    for chunk in re.split(r"\s+", text):
        if not chunk:
            continue
        for term in chunk.split("|"):
            term = term.strip()
            if term:
                candidates.append(term)

    cleaned_terms: list[str] = []
    for term in candidates:
        if re.fullmatch(r"[A-Za-z0-9]+", term):
            continue
        if re.search(r"[A-Za-z0-9]", term):
            term = re.sub(r"[A-Za-z0-9]", "", term)
            term = term.strip("()（）[]【】「」『』")
        term = term.strip()
        if term:
            cleaned_terms.append(term)
    return cleaned_terms


def load_terms_from_xlsx(
        workbook_path: Path, sheet_hint: str | None = None, base_freq: int = DENOMINATOR * 10
) -> dict[str, tuple[int, str]]:
    if load_workbook is None:
        raise ModuleNotFoundError(
            "openpyxl is required to read Excel files. Install it before using --extra-xlsx."
        )
    wb = load_workbook(workbook_path, read_only=True)
    sheet_name: str
    if sheet_hint:
        if sheet_hint in wb.sheetnames:
            sheet_name = sheet_hint
        else:
            matches = [name for name in wb.sheetnames if sheet_hint in name]
            if not matches:
                raise ValueError(f"Sheet '{sheet_hint}' not found in {workbook_path}")
            sheet_name = matches[0]
    else:
        matches = [
            name for name in wb.sheetnames if "情報科全教科書用語説明付き" in name or "用語説明付き" in name
        ]
        sheet_name = matches[0] if matches else wb.sheetnames[0]

    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return {}

    try:
        term_idx = headers.index("用語")
    except ValueError as exc:
        raise ValueError(f"'用語' column not found in sheet '{sheet_name}'") from exc

    category_idx = headers.index("カテゴリ") if "カテゴリ" in headers else None

    extra_entries: dict[str, tuple[int, str]] = {}
    for row in rows:
        if row is None:
            continue
        term_cell = row[term_idx] if term_idx < len(row) else None
        if not term_cell:
            continue
        category = row[category_idx] if category_idx is not None and category_idx < len(row) else None
        tag = _normalize_category(category)
        for term in _extract_terms(term_cell):
            if term not in extra_entries:
                extra_entries[term] = (base_freq, tag)
    return extra_entries


def write_dictionary(entries: dict[str, tuple[int, str]], output_txt: Path) -> None:
    output_txt.parent.mkdir(parents=True, exist_ok=True)
    with output_txt.open("w", encoding="utf-8", newline="\n") as fd:
        for surface, (freq, pos_tag) in sorted(
                entries.items(), key=lambda item: (-item[1][0], item[0])
        ):
            fd.write(f"{surface} {freq} {pos_tag}\n")


def build_trie(entries: dict[str, tuple[int, str]], output_trie: Path) -> None:
    if datrie is None:
        raise ModuleNotFoundError(
            "datrie is required to build the trie. Install it or run with --skip-trie."
        )
    trie = datrie.Trie(string.printable)
    for surface, (freq, pos_tag) in entries.items():
        key = _encode_key(surface)
        rank = int(math.log(freq / DENOMINATOR) + 0.5)
        if key not in trie or (isinstance(trie[key], tuple) and trie[key][0] < rank):
            trie[key] = (rank, pos_tag)
        trie[_encode_rkey(surface)] = 1
    output_trie.parent.mkdir(parents=True, exist_ok=True)
    trie.save(str(output_trie))


def main() -> None:
    repo_root = Path(__file__).resolve().parents[2]
    default_input = repo_root / "mecab-ipadic"
    default_output = repo_root / "rag" / "res" / "ipadic.txt"

    parser = argparse.ArgumentParser(
        description="Generate RagTokenizer trie resources from mecab-ipadic CSV files."
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=default_input,
        help=f"Directory containing mecab-ipadic CSV files (default: {default_input})",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=default_output,
        help=f"Output dictionary text file path (default: {default_output})",
    )
    parser.add_argument(
        "--skip-trie",
        action="store_true",
        help="Skip building the .trie cache file.",
    )
    parser.add_argument(
        "--extra-xlsx",
        type=Path,
        nargs="*",
        default=[],
        help="Additional Excel files that contain a '用語' column to append to the dictionary.",
    )
    parser.add_argument(
        "--extra-sheet",
        type=str,
        default=None,
        help="Target sheet name (or substring) inside the extra Excel files.",
    )
    parser.add_argument(
        "--extra-freq",
        type=int,
        default=DENOMINATOR * 10,
        help="Frequency assigned to terms originating from the extra Excel files.",
    )
    args = parser.parse_args()

    csv_files = sorted(args.input_dir.glob("*.csv"))
    if not csv_files:
        raise FileNotFoundError(f"No CSV files found under {args.input_dir}")

    entries = load_entries(csv_files)

    for extra_path in args.extra_xlsx:
        extra_entries = load_terms_from_xlsx(
            extra_path, sheet_hint=args.extra_sheet, base_freq=args.extra_freq
        )
        for surface, candidate in extra_entries.items():
            if surface not in entries or candidate[0] > entries[surface][0]:
                entries[surface] = candidate

    output_txt = args.output if args.output.suffix == ".txt" else args.output.with_suffix(".txt")
    write_dictionary(entries, output_txt)

    if not args.skip_trie:
        output_trie = output_txt.with_suffix(output_txt.suffix + ".trie")
        build_trie(entries, output_trie)


if __name__ == "__main__":
    main()
